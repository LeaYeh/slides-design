"""Styled text helpers built on the token type scale."""
from pptx.util import Inches, Pt
from . import tokens as t

def _set_tracking(run, pts: float):
    run.font._rPr.set("spc", str(int(pts * 100)))  # spc is 1/100 pt

def textbox(slide, box, text, *, size, bold=False, color=t.INK,
            align=t.ALIGN_LEFT, upper=False, tracking=0.0, line=1.12):
    """box = (left,top,width,height) in inches. Returns the shape."""
    left, top, width, height = (Inches(v) for v in box)
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = line
    run = p.add_run()
    run.text = text.upper() if upper else text
    f = run.font
    f.name = "Helvetica Neue"
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    if tracking:
        _set_tracking(run, tracking)
    return tb

def kicker(slide, left, top, text, accent):
    """Uppercase overline in accent-800, letter-spaced."""
    return textbox(slide, (left, top, 6.0, 0.4), text,
                   size=t.TYPE["kicker"][0], bold=True,
                   color=t.ramp(accent, "c800"), upper=True,
                   tracking=t.KICKER_TRACK_PT)

def title(slide, left, top, width, text, kind="heading", color=t.INK):
    size, bold = t.TYPE[kind]
    return textbox(slide, (left, top, width, 1.6), text, size=size, bold=bold,
                   color=color, tracking=-0.4)  # slight tight tracking

def body(slide, left, top, width, text, color=t.GRAPHITE):
    return textbox(slide, (left, top, width, 3.5), text,
                   size=t.TYPE["body"][0], color=color, line=1.3)

def caption(slide, left, top, width, text):
    return textbox(slide, (left, top, width, 0.4), text,
                   size=t.TYPE["caption"][0], color=t.SLATE)
